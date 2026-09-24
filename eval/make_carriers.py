"""Carrier generator (SPEC §17 carrier matrix): wrap a text in a document format.

Every carrier is generated in code (no binaries in git). For visible carriers the text is shown
normally after the cover paragraph; for hidden carriers (HIDDEN_CARRIERS) the cover is the
visible part and the text goes into a non-visible place (comment, display:none, white 1pt, ...).

    from eval.make_carriers import make_carrier
    carrier = make_carrier("some text", "pdf_white_1pt")      # -> Carrier(data=b"%PDF...", ...)

CLI:
    python -m eval.make_carriers --text "..." --carrier pdf_white_1pt --out data/carriers/
    python -m eval.make_carriers --text "..." --all --out data/carriers/
"""
from __future__ import annotations

import argparse
import html
import io
import json
import textwrap
from dataclasses import dataclass
from email.message import EmailMessage
from email.policy import SMTP
from pathlib import Path

CARRIERS: list[str] = [
    "plain", "markdown_comment", "html_visible", "html_display_none", "html_comment",
    "pdf_visible", "pdf_white_1pt", "pdf_metadata", "docx_hidden_run", "email_html",
    "json_field", "code_comment", "png_normal", "png_low_contrast",
]
HIDDEN_CARRIERS: set[str] = {
    "markdown_comment", "html_display_none", "html_comment", "pdf_white_1pt", "pdf_metadata",
    "docx_hidden_run", "email_html", "code_comment", "png_low_contrast",
}

DEFAULT_COVER = (
    "Logistics update: from Monday the Pune warehouse moves its outbound dispatch window to "
    "06:00-14:00. Please confirm revised pickup slots with your regional coordinator by Friday. "
    "Inventory counts and delivery routes stay unchanged."
)

# name -> (extension, content type, expected channel of `text`, expected hidden_reason)
_SPECS: dict[str, tuple[str, str, str, str | None]] = {
    "plain": ("txt", "text/plain", "visible", None),
    "markdown_comment": ("md", "text/markdown", "comment", "html_comment"),
    "html_visible": ("html", "text/html", "visible", None),
    "html_display_none": ("html", "text/html", "hidden", "display_none"),
    "html_comment": ("html", "text/html", "comment", "html_comment"),
    "pdf_visible": ("pdf", "application/pdf", "visible", None),
    "pdf_white_1pt": ("pdf", "application/pdf", "hidden", "white_text"),
    "pdf_metadata": ("pdf", "application/pdf", "metadata", "pdf_metadata"),
    "docx_hidden_run": ("docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                        "hidden", "docx_vanish"),
    "email_html": ("eml", "message/rfc822", "hidden", "display_none"),
    "json_field": ("json", "application/json", "visible", None),
    "code_comment": ("py", "text/x-python", "comment", "code_comment"),
    "png_normal": ("png", "image/png", "ocr", None),
    "png_low_contrast": ("png", "image/png", "ocr_enhanced", "low_contrast"),
}

# PNG layout
PNG_WIDTH = 1400
PNG_MARGIN = 40
PNG_FONT_SIZE = 28
PNG_LINE_HEIGHT = 40
PNG_WRAP_CHARS = 50
LOW_CONTRAST_FILL = (0xF2, 0xF2, 0xF2)       # near-white on white (contrast 1.12): barely visible


@dataclass
class Carrier:
    name: str
    filename: str
    content_type: str
    data: bytes
    expected_channel: str
    expected_hidden_reason: str | None


# ---------------------------------------------------------------- generators

def _wrap(text: str, width: int) -> list[str]:
    """Wrap at spaces only (keeps e-mail addresses / hyphenated words on one line)."""
    return textwrap.wrap(text, width, break_on_hyphens=False, break_long_words=False)


def _comment_safe(text: str) -> str:
    """Keep `text` inside an HTML comment: only "-->" / "--!>" would end it early."""
    return text.replace("--!>", "-- !>").replace("-->", "-- >")


def _html_page(body: str) -> str:
    return ("<!DOCTYPE html>\n<html><head><meta charset=\"utf-8\"><title>Operations update</title></head>\n"
            f"<body>\n{body}\n</body></html>\n")


def _pdf(cover: str, text: str | None = None, *, hidden: bool = False, subject: str | None = None) -> bytes:
    import pymupdf

    doc = pymupdf.open()
    page = doc.new_page(width=612, height=792)
    y = _pdf_paragraph(doc, page, 72, cover, fontsize=11)
    if text is not None:
        if hidden:
            _pdf_paragraph(doc, page, y + 6, text, fontsize=1, color=(1, 1, 1))   # white 1pt
        else:
            _pdf_paragraph(doc, page, y + 14, text, fontsize=11)
    doc.set_metadata({"title": "Operations update", "author": "Operations Desk",
                      "subject": subject or "", "creator": "", "producer": ""})
    data = doc.tobytes(garbage=3, deflate=True)
    doc.close()
    return data


# Fonts for text Helvetica can't encode (e.g. Devanagari, Tamil): Linux (GNU FreeFont, Noto) and macOS.
_FALLBACK_FONTS = (
    "/usr/share/fonts/truetype/freefont/FreeSerif.ttf", "/usr/share/fonts/truetype/freefont/FreeSans.ttf",
    "/usr/share/fonts/truetype/noto/NotoSansDevanagari-Regular.ttf",
    "/usr/share/fonts/truetype/noto/NotoSansTamil-Regular.ttf",
    "/System/Library/Fonts/Supplemental/Arial Unicode.ttf", "/Library/Fonts/Arial Unicode.ttf",
    "/System/Library/Fonts/Supplemental/Devanagari Sangam MN.ttc", "/System/Library/Fonts/Kohinoor.ttc",
    "/System/Library/Fonts/Supplemental/Tamil Sangam MN.ttc", "/System/Library/Fonts/Supplemental/Tamil MN.ttc",
)


def _pdf_font_file(text: str) -> str | None:
    """None when base-14 Helvetica can encode `text`; else an installed font with every glyph."""
    try:
        text.encode("latin-1")
        return None
    except UnicodeEncodeError:
        pass
    import pymupdf

    needed = {ord(ch) for ch in text if not ch.isspace()}
    for path in _FALLBACK_FONTS:
        if Path(path).exists():
            font = pymupdf.Font(fontfile=path)
            if all(font.has_glyph(cp) for cp in needed):
                return path
    raise ValueError("no installed font can render this text in a PDF carrier; "
                     f"install one of: {', '.join(_FALLBACK_FONTS[:4])}")


def _pdf_paragraph(doc: "pymupdf.Document", page: "pymupdf.Page", y: float, text: str, *,  # noqa: F821
                   fontsize: float, color: tuple[float, float, float] = (0, 0, 0)) -> float:
    """Write `text` as one wrapped paragraph (one PDF text block); returns the next free y.

    insert_textbox writes nothing when the text overflows, so long texts get a smaller font.
    """
    import pymupdf

    rect = pymupdf.Rect(72, y, 612 - 72, 792 - 36)
    font_file = _pdf_font_file(text)
    font = {"fontname": "carrierfont", "fontfile": font_file} if font_file else {}
    size = fontsize
    while True:
        spare = page.insert_textbox(rect, text, fontsize=size, color=color, **font)
        if spare >= 0:
            return y + (rect.height - spare)
        if size <= 0.5:
            raise ValueError("text too long for a one-page PDF carrier")
        size = max(0.5, size * 0.8)


def _docx(cover: str, text: str) -> bytes:
    import docx

    document = docx.Document()
    document.core_properties.author = "Operations Desk"
    document.core_properties.title = "Operations update"
    document.core_properties.comments = ""               # python-docx template says "generated by python-docx"
    paragraph = document.add_paragraph(cover + " ")
    hidden = paragraph.add_run(text)
    hidden.font.hidden = True                         # <w:vanish/>
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


def _email(cover: str, text: str) -> bytes:
    msg = EmailMessage()
    msg["From"] = "Operations Desk <ops@vendor.example>"
    msg["To"] = "team@yourcompany.example"
    msg["Subject"] = "Dispatch schedule update"
    msg.set_content(cover)
    msg.add_alternative(_html_page(
        f"<p>{html.escape(cover)}</p>\n<div style=\"display:none\">{html.escape(text)}</div>"), subtype="html")
    msg.set_boundary("sentinel-carrier-boundary")
    return msg.as_bytes(policy=SMTP)


def _json(cover: str, text: str) -> bytes:
    payload = {
        "status": "ok",
        "page": 1,
        "results": [
            {"id": 101, "title": "Dispatch window update", "body": cover},
            {"id": 102, "title": "Additional note", "body": text},
        ],
    }
    return json.dumps(payload, indent=2, ensure_ascii=False).encode("utf-8")


def _code(cover: str, text: str) -> bytes:
    comment = "\n".join("# " + line for line in _wrap(text, 88)) or "#"
    source = (
        "import datetime\n\n\n"
        "def dispatch_window(day):\n"
        f"{textwrap.indent(comment, '    ')}\n"
        "    start = datetime.time(6, 0)\n"
        "    end = datetime.time(14, 0)\n"
        "    return start, end\n"
    )
    return source.encode("utf-8")


def _png(cover: str, text: str | None, *, text_fill: tuple[int, int, int]) -> bytes:
    from PIL import Image, ImageDraw, ImageFont

    font = ImageFont.load_default(size=PNG_FONT_SIZE)
    cover_lines = _wrap(cover, PNG_WRAP_CHARS)
    text_lines = _wrap(text, PNG_WRAP_CHARS) if text else []
    rows = len(cover_lines) + (1 + len(text_lines) if text_lines else 0)
    height = 2 * PNG_MARGIN + rows * PNG_LINE_HEIGHT
    image = Image.new("RGB", (PNG_WIDTH, height), (255, 255, 255))
    draw = ImageDraw.Draw(image)
    y = PNG_MARGIN
    for line in cover_lines:
        draw.text((PNG_MARGIN, y), line, fill=(0, 0, 0), font=font)
        y += PNG_LINE_HEIGHT
    y += PNG_LINE_HEIGHT if text_lines else 0
    for line in text_lines:
        draw.text((PNG_MARGIN, y), line, fill=text_fill, font=font)
        y += PNG_LINE_HEIGHT
    buf = io.BytesIO()
    image.save(buf, format="PNG")
    return buf.getvalue()


def _build(name: str, text: str, cover: str) -> bytes:
    esc_cover, esc_text = html.escape(cover), html.escape(text)
    if name == "plain":
        return f"{cover}\n\n{text}\n".encode("utf-8")
    if name == "markdown_comment":
        return f"# Operations update\n\n{cover}\n\n<!-- {_comment_safe(text)} -->\n".encode("utf-8")
    if name == "html_visible":
        return _html_page(f"<p>{esc_cover}</p>\n<p>{esc_text}</p>").encode("utf-8")
    if name == "html_display_none":
        return _html_page(f"<p>{esc_cover}</p>\n<div style=\"display:none\">{esc_text}</div>").encode("utf-8")
    if name == "html_comment":
        return _html_page(f"<p>{esc_cover}</p>\n<!-- {_comment_safe(text)} -->").encode("utf-8")
    if name == "pdf_visible":
        return _pdf(cover, text)
    if name == "pdf_white_1pt":
        return _pdf(cover, text, hidden=True)
    if name == "pdf_metadata":
        return _pdf(cover, subject=text)
    if name == "docx_hidden_run":
        return _docx(cover, text)
    if name == "email_html":
        return _email(cover, text)
    if name == "json_field":
        return _json(cover, text)
    if name == "code_comment":
        return _code(cover, text)
    if name == "png_normal":
        return _png(cover, text, text_fill=(0, 0, 0))
    if name == "png_low_contrast":
        return _png(cover, text, text_fill=LOW_CONTRAST_FILL)
    raise ValueError(f"unknown carrier {name!r}; choose from {', '.join(CARRIERS)}")


def make_carrier(text: str, carrier: str, cover: str | None = None) -> Carrier:
    """Embed `text` in the named carrier; `cover` is the benign visible text (default: DEFAULT_COVER)."""
    if carrier not in _SPECS:
        raise ValueError(f"unknown carrier {carrier!r}; choose from {', '.join(CARRIERS)}")
    ext, content_type, channel, reason = _SPECS[carrier]
    data = _build(carrier, text, cover if cover is not None else DEFAULT_COVER)
    return Carrier(name=carrier, filename=f"{carrier}.{ext}", content_type=content_type, data=data,
                   expected_channel=channel, expected_hidden_reason=reason)


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Write carrier documents that embed a text.")
    parser.add_argument("--text", required=True, help="text to embed (attack or benign)")
    group = parser.add_mutually_exclusive_group(required=True)
    group.add_argument("--carrier", choices=CARRIERS, help="one carrier")
    group.add_argument("--all", action="store_true", help="write every carrier")
    parser.add_argument("--cover", default=None, help="benign visible cover text")
    parser.add_argument("--out", default="data/carriers", help="output directory")
    args = parser.parse_args(argv)
    out_dir = Path(args.out)
    out_dir.mkdir(parents=True, exist_ok=True)
    for name in (CARRIERS if args.all else [args.carrier]):
        carrier = make_carrier(args.text, name, args.cover)
        path = out_dir / carrier.filename
        path.write_bytes(carrier.data)
        reason = carrier.expected_hidden_reason
        where = carrier.expected_channel + (f"/{reason}" if reason else "")
        print(f"{path}  ({where})")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
